"""Import craft villages from the data-collection workbook (docs/data/craft-villages.xlsx).

Sheet layout: row 1 holds Vietnamese labels for humans, row 2 holds the field keys used
here, and the data starts on row 3. Villages are matched by ``code``, so re-importing an
updated workbook refreshes the existing rows instead of duplicating them.
"""

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.db import transaction

from apps.core.models import PublicationStatus

from .models import CraftType, CraftVillage, FormerRegion, RecognitionStatus

SHEET_NAME = "Villages"
KEY_ROW = 2
FIRST_DATA_ROW = 3

# The workbook is filled in Vietnamese; map those labels onto the stored choice values.
CRAFT_TYPE_BY_LABEL = {
    "nón lá": CraftType.CONICAL_HAT,
    "nước mắm": CraftType.FISH_SAUCE,
    "dệt chiếu": CraftType.MAT_WEAVING,
    "mây tre đan": CraftType.BAMBOO_RATTAN,
    "mộc - chạm khắc": CraftType.WOODWORK,
    "rượu": CraftType.RICE_WINE,
    "bún - bánh": CraftType.NOODLES_CAKES,
    "rèn - đúc": CraftType.METALWORK,
    "quạt giấy": CraftType.PAPER_FAN,
    "chổi đót": CraftType.BROOM,
    "thực phẩm - đặc sản khác": CraftType.FOOD_SPECIALTY,
    "khác": CraftType.OTHER,
}
RECOGNITION_BY_LABEL = {
    "làng nghề truyền thống": RecognitionStatus.TRADITIONAL_VILLAGE,
    "làng nghề": RecognitionStatus.VILLAGE,
    "nghề truyền thống": RecognitionStatus.TRADITIONAL_CRAFT,
    "chưa công nhận": RecognitionStatus.NOT_RECOGNIZED,
    "chưa rõ": RecognitionStatus.UNKNOWN,
}
REGION_BY_LABEL = {
    "quảng trị cũ": FormerRegion.QUANG_TRI,
    "quảng bình cũ": FormerRegion.QUANG_BINH,
}

# Workbook field key -> model field name (only where the two differ).
FIELD_ALIASES = {
    "commune_new": "commune",
    "commune_old": "former_commune",
    "district_old": "former_district",
    "current_status": "current_situation",
}
TEXT_FIELDS = [
    "name",
    "main_products",
    "recognition_decision",
    "hamlet",
    "commune",
    "former_commune",
    "former_district",
    "short_description",
    "history",
    "production_process",
    "current_situation",
    "culture",
    "visitor_info",
    "contact",
    "sources",
    "notes",
]
REQUIRED_FIELDS = ["code", "name", "craft_type", "former_region"]


@dataclass
class ImportReport:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)

    def summary(self) -> str:
        return (
            f"created={self.created} updated={self.updated} "
            f"skipped={self.skipped} errors={len(self.errors)}"
        )


def _clean(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _choice(value: Any, mapping: dict[str, Any], row: int, column: str, report: ImportReport):
    label = _clean(value).lower()
    if not label:
        return None
    if label not in mapping:
        report.errors.append(f"Row {row}: unknown {column} '{_clean(value)}'")
        return None
    return mapping[label]


def _decimal(value: Any, row: int, column: str, report: ImportReport) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError):
        report.errors.append(f"Row {row}: invalid {column} '{value}'")
        return None


def _year(value: Any, row: int, report: ImportReport) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        report.errors.append(f"Row {row}: invalid recognition_year '{value}'")
        return None


def import_villages(path: Path, *, dry_run: bool = False, publish: bool = False) -> ImportReport:
    from openpyxl import load_workbook  # imported lazily: only needed for imports

    report = ImportReport()
    workbook = load_workbook(path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Workbook has no '{SHEET_NAME}' sheet")
    sheet = workbook[SHEET_NAME]

    keys = [_clean(cell.value) for cell in sheet[KEY_ROW]]
    missing = [f for f in REQUIRED_FIELDS if f not in keys and FIELD_ALIASES.get(f, f) not in keys]
    if missing:
        raise ValueError(f"Workbook is missing required columns: {', '.join(missing)}")

    with transaction.atomic():
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), start=FIRST_DATA_ROW
        ):
            raw = {
                FIELD_ALIASES.get(key, key): value
                for key, value in zip(keys, row, strict=False)
                if key
            }
            code = _clean(raw.get("code"))
            if not code:
                continue

            craft_type = _choice(
                raw.get("craft_type"), CRAFT_TYPE_BY_LABEL, row_number, "craft_type", report
            )
            region = _choice(
                raw.get("former_region"), REGION_BY_LABEL, row_number, "former_region", report
            )
            name = _clean(raw.get("name"))
            if not (name and craft_type and region):
                report.skipped += 1
                if name and (craft_type is None or region is None):
                    continue
                report.errors.append(f"Row {row_number}: missing name, craft type or region")
                continue

            values = {f: _clean(raw.get(f)) for f in TEXT_FIELDS}
            values.update(
                craft_type=craft_type,
                former_region=region,
                recognition_status=_choice(
                    raw.get("recognition_status"),
                    RECOGNITION_BY_LABEL,
                    row_number,
                    "recognition_status",
                    report,
                )
                or RecognitionStatus.UNKNOWN,
                recognition_year=_year(raw.get("recognition_year"), row_number, report),
                latitude=_decimal(raw.get("latitude"), row_number, "latitude", report),
                longitude=_decimal(raw.get("longitude"), row_number, "longitude", report),
            )
            if publish:
                # Without --publish the editorial status is left untouched, so re-importing
                # never unpublishes a village an editor already reviewed.
                values["status"] = PublicationStatus.PUBLISHED

            _, created = CraftVillage.objects.update_or_create(code=code, defaults=values)
            if created:
                report.created += 1
            else:
                report.updated += 1

        if dry_run:
            transaction.set_rollback(True)

    workbook.close()
    return report
